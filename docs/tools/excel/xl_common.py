from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

F = "Arial"
HEAD = PatternFill("solid", fgColor="0F6CBD")
INPUT = PatternFill("solid", fgColor="FFF7CC")
thin = Side(style="thin", color="BFC7D5")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

def header(ws, cols, row=1, freeze_col=2):
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=freeze_col)

def write_rows(ws, rows, start=2, input_cols=(), italic=False):
    for r, values in enumerate(rows, start):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=F, size=10, italic=italic)
            cell.alignment = WRAP
            cell.border = BOX
            if c in input_cols:
                cell.fill = INPUT

def dropdown(ws, rng, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)

def title(ws, cell, text):
    ws[cell] = text
    ws[cell].font = Font(name=F, bold=True, size=14)
