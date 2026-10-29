import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font
from backlog_data import BACKLOG, SESSIONS, SPRINTS, WORKLOG
from xl_common import F, dropdown, header, title, write_rows

STATUSES = ["Hoàn thành", "Đang làm", "Chưa làm"]
wb = Workbook()

sm = wb.active
sm.title = "Tổng quan"
title(sm, "A1", "Quản lý dự án — Hệ thống theo dõi bệnh nhân (CDTN2)")
sm.column_dimensions["A"].width = 34
sm.column_dimensions["B"].width = 14
sm.column_dimensions["C"].width = 14
n = len(BACKLOG) + 1
pts = f"'Product backlog'!$F$2:$F${n + 50}"
st = f"'Product backlog'!$H$2:$H${n + 50}"
sp = f"'Product backlog'!$G$2:$G${n + 50}"
sm["A3"], sm["B3"], sm["C3"] = "Trạng thái", "Số mục", "Story point"
for i, s in enumerate(STATUSES, 4):
    sm.cell(row=i, column=1, value=s)
    sm.cell(row=i, column=2, value=f'=COUNTIF({st},"{s}")')
    sm.cell(row=i, column=3, value=f'=SUMIF({st},"{s}",{pts})')
sm["A7"], sm["B7"], sm["C7"] = "Tổng", "=SUM(B4:B6)", "=SUM(C4:C6)"
sm["A8"], sm["C8"] = "Tiến độ theo story point", "=IFERROR(C4/C7,0)"
sm["C8"].number_format = "0.0%"
sm["A10"], sm["B10"], sm["C10"] = "Sprint", "Số mục", "Story point"
for i, name in enumerate(["Sprint 1", "Sprint 1-2", "Sprint 2", "Sprint 3", "Sprint 4"], 11):
    sm.cell(row=i, column=1, value=name)
    sm.cell(row=i, column=2, value=f'=COUNTIF({sp},"{name}")')
    sm.cell(row=i, column=3, value=f'=SUMIF({sp},"{name}",{pts})')
for row in sm.iter_rows(min_row=3, max_row=13):
    for c in row:
        c.font = Font(name=F, bold=c.row in (3, 7, 10))
notes = ["Cách dùng:",
         "• Ô nền vàng là ô nhóm cần cập nhật (người phụ trách, trạng thái, nội dung nhật ký...).",
         "• Trạng thái chọn từ danh sách; bảng tổng quan tự tính lại.",
         "• Nhật ký buổi học là minh chứng chấm Rubric 2 (30%): cập nhật trong tuần, không dồn cuối kỳ.",
         "• Nhật ký công việc: mỗi dòng một việc, ghi mã commit để đối chiếu với GitHub."]
for i, t in enumerate(notes, 16):
    sm.cell(row=i, column=1, value=t).font = Font(name=F, bold=(i == 16))

pb = wb.create_sheet("Product backlog")
header(pb, [("ID", 7), ("Yêu cầu", 9), ("User story", 46), ("Tiêu chí chấp nhận", 44), ("Ưu tiên", 11),
            ("Story point", 8), ("Sprint", 11), ("Trạng thái", 12), ("Test case", 16), ("Người phụ trách", 16)], freeze_col=3)
write_rows(pb, [list(r) + [""] for r in BACKLOG], input_cols={8, 10})
dropdown(pb, f"H2:H{n + 50}", STATUSES)
dropdown(pb, f"E2:E{n + 50}", ["Cao", "Trung bình", "Thấp"])
pb.auto_filter.ref = f"A1:J{n}"

sp_ws = wb.create_sheet("Kế hoạch sprint")
header(sp_ws, [("Sprint", 12), ("Thời gian", 18), ("Mục tiêu", 50), ("Hạng mục backlog", 28), ("Trạng thái", 22)])
write_rows(sp_ws, SPRINTS, input_cols={5})

nk = wb.create_sheet("Nhật ký buổi học")
header(nk, [("Buổi", 6), ("Ngày", 12), ("Nội dung buổi", 26), ("Bàn giao chính", 32), ("Công việc đã làm (ai làm gì)", 40),
            ("Sản phẩm bàn giao", 30), ("Khó khăn / quyết định", 30), ("Góp ý của GVHD", 30), ("Kế hoạch tuần tới", 30)], freeze_col=3)
write_rows(nk, [[b, "", c, d, "", "", "", "", ""] for b, c, d in SESSIONS], input_cols={2, 5, 6, 7, 8, 9})

wl = wb.create_sheet("Nhật ký công việc")
header(wl, [("Ngày", 12), ("Thành viên", 16), ("Công việc", 56), ("Kết quả", 40), ("Commit", 11)], freeze_col=3)
write_rows(wl, WORKLOG, input_cols={1, 2, 3, 4, 5})

tv = wb.create_sheet("Thành viên")
header(tv, [("Họ tên", 24), ("MSSV", 14), ("Vai trò trong nhóm", 18), ("Phụ trách chính", 40)])
write_rows(tv, [["", "", "Nhóm trưởng", "Backend / API"], ["", "", "Thành viên", "Web nhân viên + quản trị"],
                ["", "", "Thành viên", "PWA người nhà + thông báo"], ["", "", "Thành viên", "Kiểm thử + tài liệu"]], input_cols={1, 2, 3, 4})

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "quan-ly", "Quan-ly-du-an.xlsx")
wb.save(OUT)
print("saved", OUT)
