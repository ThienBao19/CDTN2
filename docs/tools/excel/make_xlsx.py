import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from bug_data import BUGS
from tc_data import ROWS as BASE_ROWS
from tc_extra import EXTRA

from tc_v2 import EXTRA2, REPLACE, SHARED_V2

ROWS = [REPLACE.get(r[0], r) for r in BASE_ROWS + EXTRA] + EXTRA2

# Ghi chú cho lỗi gắn với luồng mã/QR cũ và rủi ro của cách vào xem mới
_NOTE = " — Luồng mã/QR đã thay bằng nhân viên đăng ký người nhà (17/10/2026)"
BUGS = [b[:9] + (b[9] + _NOTE,) if b[0] in ("BUG-02", "BUG-06") else b for b in BUGS] + [
    ("HC-07", "TC-03, TC-05", "Người nhà vào xem không cần mật khẩu: ai biết đúng SĐT, họ tên người nhà và họ tên bệnh nhân đã được đăng ký thì xem được",
     "Nhập đúng 3 thông tin của một người nhà đã được đăng ký", "Trung bình", "Hạn chế đã biết", "", "17/10/2026", "",
     "Theo yêu cầu (không mật khẩu, không xác nhận). Giảm thiểu: chỉ người nhân viên đăng ký mới vào được, khóa 5 lần sai/15 phút theo SĐT, phiên hết khi ca kết thúc, nhật ký FAMILY_ACCESS"),
]

DOCS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(DOCS, "kiem-thu", "Ket-qua-kiem-thu.xlsx")
F = "Arial"
HEAD = PatternFill("solid", fgColor="0F6CBD")
INPUT = PatternFill("solid", fgColor="FFF7CC")
thin = Side(style="thin", color="BFC7D5")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

def header(ws, cols, row=1):
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=3)

import os
ANH = os.path.join(DOCS, "kiem-thu", "anh")
FILES = sorted(os.listdir(ANH)) if os.path.isdir(ANH) else []

# Ảnh minh chứng: ảnh giao diện theo mã TC; TC chỉ kiểm bằng test tự động dùng ảnh kết quả chạy test
# Ngày test của các test case thêm sau đợt đầu
TEST_DATES = {"TC-31": "29/10/2026", "TC-52": "29/10/2026", "TC-53": "29/10/2026", "TC-54": "29/10/2026"}


def evidence_for(code, kind, is_done):
    if not is_done:
        return ""
    if code in SHARED_V2:
        return SHARED_V2[code]
    own = [f for f in FILES if f == code + ".png" or f.startswith(code + "a-")]
    if own:
        return ", ".join("anh/" + f for f in own)
    shared = {
        "TC-21": "anh/TC-26.png",
        "TC-46": "anh/TC-24a-sua-ca.png",
        "TC-42": "anh/KQ-07-sao-luu-khoi-phuc.png",
        "TC-48": "anh/KQ-07-sao-luu-khoi-phuc.png",
        "TC-43": "anh/KQ-06-trien-khai-docker.png",
    }
    if code in shared:
        return shared[code]
    if code == "TC-33":
        return "anh/KQ-05-k6-500-nguoi-nha.png"
    if "k6" in kind:
        return "anh/KQ-04-k6-100-nguoi-nha.png"
    if "e2e" in kind:
        return "anh/KQ-02-e2e-test.png"
    if "Unit" in kind:
        return "anh/KQ-01-unit-test.png"
    return ""

wb = Workbook()
ws = wb.active
ws.title = "Test case"
cols = [("Mã TC", 8), ("Yêu cầu", 13), ("Kịch bản", 30), ("Bước thực hiện / dữ liệu", 36),
        ("Kết quả mong đợi", 34), ("Loại kiểm thử", 15), ("Test tự động (file › tên)", 34),
        ("Kết quả thực tế", 34), ("Trạng thái", 11), ("Ảnh minh chứng", 20),
        ("Người test", 14), ("Ngày test", 12)]
header(ws, cols)
INPUT_COLS = {8, 9, 10, 11, 12}
for r, row in enumerate(ROWS, 2):
    code, fr, scen, steps, expect, kind, auto, actual, status = row
    is_done = status == "Đạt"
    values = [code, fr, scen, steps, expect, kind, auto, actual, status,
              evidence_for(code, kind, is_done), "", TEST_DATES.get(code, "15/10/2026") if is_done else ""]
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=F, size=10)
        cell.alignment = WRAP
        cell.border = BOX
        if c in INPUT_COLS:
            cell.fill = INPUT
last = len(ROWS) + 1
dv = DataValidation(type="list", formula1='"Đạt,Không đạt,Chưa test"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"I2:I{last + 50}")
ws.auto_filter.ref = f"A1:L{last}"

bug = wb.create_sheet("Bug log")
header(bug, [("Mã lỗi", 9), ("Mã TC", 9), ("Mô tả lỗi", 40), ("Bước tái hiện", 36),
             ("Mức độ", 12), ("Trạng thái", 13), ("Người sửa", 14), ("Ngày phát hiện", 14),
             ("Ngày sửa", 12), ("Ghi chú / hạn chế đã biết", 34)])
write_bug_rows = True
for r, values in enumerate(BUGS, 2):
    for c, v in enumerate(values, 1):
        cell = bug.cell(row=r, column=c, value=v)
        cell.font = Font(name=F, size=10)
        cell.alignment = WRAP
        cell.border = BOX
for name, opts in [("E", "Cao,Trung bình,Thấp"), ("F", "Mới,Đang sửa,Đã sửa,Hạn chế đã biết")]:
    v = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
    bug.add_data_validation(v)
    v.add(f"{name}2:{name}200")

sm = wb.create_sheet("Tổng hợp", 0)
sm.column_dimensions["A"].width = 30
sm.column_dimensions["B"].width = 14
sm["A1"] = "Kết quả kiểm thử — Hệ thống theo dõi bệnh nhân"
sm["A1"].font = Font(name=F, bold=True, size=14)
rng = f"'Test case'!$I$2:$I${last + 50}"
lines = [("Tổng số test case", f"=COUNTA('Test case'!$A$2:$A${last + 50})"),
         ("Đạt", f'=COUNTIF({rng},"Đạt")'),
         ("Không đạt", f'=COUNTIF({rng},"Không đạt")'),
         ("Chưa test", f'=COUNTIF({rng},"Chưa test")'),
         ("Tỷ lệ đạt (trên số đã test)", "=IFERROR(B4/(B4+B5),0)"),
         ("Số lỗi đang mở", "=COUNTIF('Bug log'!$F$2:$F$200,\"Mới\")+COUNTIF('Bug log'!$F$2:$F$200,\"Đang sửa\")")]
for i, (label, formula) in enumerate(lines, 3):
    sm.cell(row=i, column=1, value=label).font = Font(name=F)
    c = sm.cell(row=i, column=2, value=formula)
    c.font = Font(name=F, bold=True)
sm["B7"].number_format = "0.0%"
notes = [
    "Cách dùng:",
    "• Ô nền vàng ở sheet 'Test case' là ô cần điền: kết quả thực tế, trạng thái, ảnh, người test, ngày.",
    "• Cột 'Trạng thái' chọn từ danh sách: Đạt / Không đạt / Chưa test. Bảng này tự đếm lại.",
    "• Ảnh minh chứng lưu ở docs/kiem-thu/anh/, đặt tên theo mã TC (VD anh/TC-01.png).",
    "  Mỗi TC chụp: (1) màn hình giao diện hoặc Postman, (2) terminal kết quả test tự động.",
    "• Chạy test tự động: pnpm --filter @pt/api test (unit) và pnpm --filter @pt/api test:e2e (API).",
    "• Lỗi phát hiện ghi ở sheet 'Bug log'; dòng 2 là ví dụ cách ghi.",
    "• Kết quả ngày 15/10/2026 lấy từ lần chạy: 55 unit + 56 e2e (Jest) + 22 kịch bản giao diện (Playwright) đều đạt; coverage gộp 94,5% statements (cập nhật 17/10/2026 sau khi đổi cách người nhà vào xem; 29/10/2026 thêm âm báo, chuông kết thúc, giữ phiên khi mở lại app).",
]
for i, text in enumerate(notes, 11):
    sm.cell(row=i, column=1, value=text).font = Font(name=F, bold=(i == 11))

wb.save(OUT)
print("saved", OUT, "rows", len(ROWS))
